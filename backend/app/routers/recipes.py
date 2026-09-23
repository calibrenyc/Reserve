from decimal import Decimal
from typing import List, Optional
import io
import json
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Request
from sqlalchemy.orm import Session

from backend.app.database import get_db, current_location_id
from backend.app.models import Recipe, RecipeVariant, RecipeIngredient, InventoryItem, RecipeImport
from backend.app.schemas import RecipeCreate
from backend.app.services.avt_engine import AvTEngine

router = APIRouter(prefix="/api/recipes", tags=["Recipes"])

def _header_key(value):
    return " ".join("".join(char.lower() if char.isalnum() else " " for char in str(value or "")).split())

import logging
from fastapi.responses import JSONResponse

logger = logging.getLogger("reserve.recipes_import")

STANDARD_RECIPE_COLUMNS = {
    "recipe_name": {"recipe name", "recipe"}, "variant": {"variant", "size", "recipe variant"},
    "recipe_type": {"recipe type", "type"}, "area": {"area", "station"}, "category": {"category"},
    "yield_quantity": {"yield qty", "yield quantity", "yield"}, "yield_unit": {"yield unit"},
    "ingredient": {"ingredient", "ingredient name", "item", "primary ingredients", "primary ingredient"}, "quantity": {"quantity", "qty", "amount"}, "unit": {"unit", "uom", "units"},
    "notes": {"notes", "ingredient notes"},
}

@router.post("/imports/preview")
async def preview_recipe_import(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Store a deterministic XLSX import review; no recipe data is committed here."""
    logger.info(f"[Recipe Import Request] URL={request.url.path} Filename={file.filename}")
    stage = "INITIALIZATION"
    importer = "unknown"
    sheet_names = []
    
    if not (file.filename or "").lower().endswith(".xlsx"):
        return JSONResponse(
            status_code=400,
            content={
                "success": False,
                "error": {
                    "code": "RECIPE_IMPORT_INVALID_FILE",
                    "message": "Upload an .xlsx recipe workbook."
                },
                "detail": "Upload an .xlsx recipe workbook."
            }
        )
        
    try:
        stage = "READ_XLSX_BUFFER"
        import openpyxl
        content = await file.read()
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheets = [{"name": sheet.title, "rows": max((sheet.max_row or 0) - 1, 0), "columns": sheet.max_column or 0} for sheet in workbook.worksheets]
        sheet_names = [sheet.title for sheet in workbook.worksheets]
        
        stage = "DETECT_TEMPLATE"
        candidate = next((sheet for sheet in workbook.worksheets if sheet.title.strip().casefold() in {"recipe template", "recipes"}), None)
        candidate = candidate or next((sheet for sheet in workbook.worksheets if "recipe" in sheet.title.casefold()), None)
        extracted = []
        header_index = None
        normalized_sheet_names = {sheet.title.casefold(): sheet for sheet in workbook.worksheets}
        
        variant_count = 0
        if {"recipes", "recipe variants", "recipe ingredients"}.issubset(normalized_sheet_names):
            stage = "PARSE_RESERVE_NORMALIZED_RECIPES"
            importer = "reserve_normalized_recipes"
            def records(sheet):
                data = list(sheet.iter_rows(values_only=True))
                if not data:
                    return []
                headers = [_header_key(value) for value in data[0]]
                return [dict(zip(headers, row)) for row in data[1:] if any(value not in (None, "") for value in row)]
            recipe_rows = records(normalized_sheet_names["recipes"])
            variant_rows = records(normalized_sheet_names["recipe variants"])
            ingredient_rows = records(normalized_sheet_names["recipe ingredients"])
            
            variants = {str(row.get("variant key")): row for row in variant_rows}
            recipes = {str(row.get("recipe key")): row for row in recipe_rows}
            variant_count = len(variants)
            
            org_id = getattr(request.state, "organization_id", None)
            items = db.query(InventoryItem).execution_options(skip_tenant_scope=True).filter(InventoryItem.organization_id == org_id).all() if org_id else []
            by_name = {_header_key(item.name): item for item in items}
            
            for row in ingredient_rows:
                variant = variants.get(str(row.get("variant key")))
                recipe = recipes.get(str(variant.get("recipe key"))) if variant else None
                if not variant or not recipe: continue
                raw_name = str(row.get("ingredient name") or "").strip()
                matched = by_name.get(_header_key(raw_name))
                extracted.append({
                    "recipe_key": recipe.get("recipe key"),
                    "recipe_name": recipe.get("name"),
                    "variant_key": row.get("variant key"),
                    "variant": variant.get("variant name") or "Standard",
                    "recipe_type": recipe.get("recipe type") or "MENU_ITEM",
                    "area": recipe.get("area"),
                    "category": recipe.get("category"),
                    "yield_quantity": variant.get("yield quantity") or 1,
                    "yield_unit": variant.get("yield unit") or "EA",
                    "ingredient": raw_name,
                    "quantity": row.get("quantity"),
                    "unit": row.get("unit"),
                    "reference_type": row.get("reference type") or "ITEM",
                    "item_id": matched.id if matched else None,
                    "match_status": "EXACT" if matched else "UNRESOLVED",
                    "notes": row.get("notes"),
                    "sort_order": row.get("sort order") or 0
                })
        elif candidate:
            stage = "PARSE_CANDIDATE_SHEET"
            sheet_rows = [row for row in candidate.iter_rows(values_only=True) if any(value not in (None, "") for value in row)]
            header_index = next((index for index, row in enumerate(sheet_rows[:30]) if any(_header_key(value) in STANDARD_RECIPE_COLUMNS["recipe_name"] for value in row) and any(_header_key(value) in STANDARD_RECIPE_COLUMNS["ingredient"] for value in row)), None)
            if header_index is None and sheet_rows:
                header_index = next((index for index, row in enumerate(sheet_rows[:30]) if sum(_header_key(value) in aliases for aliases in STANDARD_RECIPE_COLUMNS.values() for value in row) >= 2), 0)
            headers = sheet_rows[header_index] if header_index is not None and header_index < len(sheet_rows) else ()
            rows = iter(sheet_rows[header_index + 1:]) if header_index is not None and header_index < len(sheet_rows) else iter([])
            mapping = {field: next((index for index, value in enumerate(headers) if _header_key(value) in aliases), None) for field, aliases in STANDARD_RECIPE_COLUMNS.items()}
            if mapping["recipe_name"] is not None and mapping["ingredient"] is not None:
                importer = "reserve_standard_recipes"
                org_id = getattr(request.state, "organization_id", None)
                items = db.query(InventoryItem).execution_options(skip_tenant_scope=True).filter(InventoryItem.organization_id == org_id).all() if org_id else []
                by_name = {_header_key(item.name): item for item in items}
                for row in rows:
                    value = lambda field: str(row[mapping[field]] or "").strip() if mapping.get(field) is not None and mapping[field] < len(row) else ""
                    if not value("recipe_name") or not value("ingredient"): continue
                    matched = by_name.get(_header_key(value("ingredient")))
                    extracted.append({"recipe_name": value("recipe_name"), "variant": value("variant") or "Standard", "recipe_type": value("recipe_type") or "MENU_ITEM", "area": value("area"), "ingredient": value("ingredient"), "quantity": value("quantity"), "unit": value("unit"), "item_id": matched.id if matched else None, "match_status": "EXACT" if matched else "UNRESOLVED", "notes": value("notes")})
                variant_count = len({row["variant"] for row in extracted})
            else:
                importer = "solbol_inventory_tracker"

        stage = "BUILD_REVIEW_SUMMARY"
        recipe_count = len({row["recipe_name"] for row in extracted})
        unresolved = sum(row["match_status"] == "UNRESOLVED" for row in extracted)
        
        warnings_list = []
        if unresolved:
            warnings_list.append(f"{unresolved} ingredients need mapping.")
        elif not candidate and importer != "reserve_normalized_recipes":
            warnings_list.append("No recipe sheet was found.")
            
        if "import warnings" in normalized_sheet_names:
            warn_sheet = normalized_sheet_names["import warnings"]
            warn_rows = [row for row in warn_sheet.iter_rows(values_only=True) if any(value not in (None, "") for value in row)]
            if len(warn_rows) > 1:
                for w_row in warn_rows[1:]:
                    w_txt = " - ".join(str(v).strip() for v in w_row if v not in (None, ""))
                    if w_txt and w_txt not in warnings_list:
                        warnings_list.append(w_txt)

        review = {
            "sheets": sheets,
            "detected_template": importer,
            "recipe_sheet": candidate.title if candidate else (normalized_sheet_names["recipes"].title if "recipes" in normalized_sheet_names else None),
            "recipes": extracted,
            "recipes_found": recipe_count,
            "recipesDetected": recipe_count,
            "variantsDetected": variant_count,
            "ingredientLinesDetected": len(extracted),
            "ingredient_lines": len(extracted),
            "unresolved": unresolved,
            "status": "Needs Review",
            "message": f"Workbook uploaded ({importer}). Review matches before importing.",
            "header_row": header_index + 1 if header_index is not None else None
        }

        stage = "RECORD_DATABASE_LOG"
        record = RecipeImport(
            organization_id=getattr(request.state, "organization_id", None),
            location_id=getattr(request.state, "location_id", None),
            filename=file.filename or "recipe-workbook.xlsx",
            importer=review["detected_template"],
            uploaded_by=getattr(request.state, "user_id", None),
            recipes_detected=recipe_count,
            warnings_json=json.dumps(warnings_list),
            review_json=json.dumps(review)
        )
        db.add(record)
        db.commit()
        db.refresh(record)

        logger.info(f"[Recipe Import Success] URL={request.url.path} Importer={importer} Recipes={recipe_count} Variants={variant_count} Ingredients={len(extracted)}")
        return {
            "success": True,
            "import_id": record.id,
            "detected_template": importer,
            "recipesDetected": recipe_count,
            "variantsDetected": variant_count,
            "ingredientLinesDetected": len(extracted),
            "warnings": warnings_list,
            **review
        }
    except Exception as exc:
        logger.error(
            f"[Recipe Import Exception] URL={request.url.path} Filename={file.filename} "
            f"HTTPStatus=400 ContentType=application/json Stage={stage} Sheets={sheet_names} Importer={importer} Exception={exc}",
            exc_info=True
        )
        return JSONResponse(
            status_code=400,
            content={
                "success": False,
                "error": {
                    "code": "RECIPE_IMPORT_PARSE_FAILED",
                    "message": f"Unable to parse recipe spreadsheet ({stage}): {str(exc)}"
                },
                "detail": f"Unable to parse recipe spreadsheet ({stage}): {str(exc)}"
            }
        )

def _find_item_match(raw_name: str, items: List[InventoryItem]) -> Optional[InventoryItem]:
    key = _header_key(raw_name)
    if not key: return None
    by_name = {_header_key(item.name): item for item in items}
    if key in by_name:
        return by_name[key]
    words = [w for w in key.split() if len(w) > 2]
    best_item = None
    best_score = 0
    for item in items:
        item_key = _header_key(item.name)
        if key in item_key or item_key in key:
            if item.current_cost and Decimal(str(item.current_cost)) > 0:
                return item
            best_item = best_item or item
        item_words = set(item_key.split())
        overlap = sum(1 for w in words if w in item_words)
        if overlap > best_score:
            best_score = overlap
            best_item = item
    return best_item if best_score > 0 else None

@router.post("/imports/{import_id}/commit")
def commit_recipe_import(import_id: str, request: Request, db: Session = Depends(get_db)):
    """Commit previewed recipe import rows into active Recipe, RecipeVariant, and RecipeIngredient database models."""
    import_record = db.query(RecipeImport).filter(RecipeImport.id == import_id).first()
    if not import_record:
        raise HTTPException(404, "Recipe import preview record not found.")

    try:
        review_data = json.loads(import_record.review_json or "{}")
        extracted_recipes = review_data.get("recipes", [])
        if not extracted_recipes:
            raise HTTPException(400, "No recipes found in this workbook to commit.")

        org_id = getattr(request.state, "organization_id", None)
        loc_id = getattr(request.state, "location_id", None) or current_location_id.get()
        if not loc_id:
            loc = db.query(Location).filter(Location.organization_id == org_id).first() if org_id else db.query(Location).first()
            if loc:
                loc_id = loc.id
                if not org_id:
                    org_id = loc.organization_id

        items = db.query(InventoryItem).execution_options(skip_tenant_scope=True).filter(InventoryItem.organization_id == org_id).all() if org_id else db.query(InventoryItem).all()

        created_recipe_ids = set()
        grouped_recipes = {}
        for row in extracted_recipes:
            r_name = (row.get("recipe_name") or "").strip()
            if not r_name: continue
            if r_name not in grouped_recipes:
                grouped_recipes[r_name] = []
            grouped_recipes[r_name].append(row)

        for r_name, rows in grouped_recipes.items():
            first_row = rows[0]
            recipe = db.query(Recipe).filter(Recipe.name == r_name, Recipe.location_id == loc_id).first()
            if not recipe:
                recipe = Recipe(
                    name=r_name,
                    organization_id=org_id,
                    location_id=loc_id,
                    category=first_row.get("category") or "Entrees",
                    area=first_row.get("area"),
                    recipe_type=first_row.get("recipe_type") or "MENU_ITEM",
                    serving_yield=Decimal(str(first_row.get("yield_quantity") or 1.0)),
                    menu_price=Decimal("12.00")
                )
                db.add(recipe)
                db.flush()

            created_recipe_ids.add(recipe.id)

            variant_groups = {}
            for row in rows:
                v_name = (row.get("variant") or "Standard").strip()
                if v_name not in variant_groups:
                    variant_groups[v_name] = []
                variant_groups[v_name].append(row)

            for v_name, v_rows in variant_groups.items():
                vf_row = v_rows[0]
                variant = db.query(RecipeVariant).filter(RecipeVariant.recipe_id == recipe.id, RecipeVariant.name == v_name).first()
                if not variant:
                    variant = RecipeVariant(
                        recipe_id=recipe.id,
                        name=v_name,
                        yield_quantity=Decimal(str(vf_row.get("yield_quantity") or 1.0)),
                        yield_unit=vf_row.get("yield_unit") or "EA",
                        serving_size=vf_row.get("serving_size")
                    )
                    db.add(variant)
                    db.flush()

                for ing_row in v_rows:
                    raw_ing_name = str(ing_row.get("ingredient") or "").strip()
                    if not raw_ing_name: continue
                    matched_item = _find_item_match(raw_ing_name, items)

                    if not matched_item:
                        default_cost = Decimal("1.75")
                        l_name = raw_ing_name.lower()
                        if any(k in l_name for k in ["cup", "lid", "straw", "bag", "napkin"]):
                            default_cost = Decimal("0.15")
                        elif any(k in l_name for k in ["milk", "juice"]):
                            default_cost = Decimal("3.50")
                        elif any(k in l_name for k in ["acai", "dragonfruit", "pitaya", "puree", "pack", "pouch"]):
                            default_cost = Decimal("1.25")

                        matched_item = InventoryItem(
                            name=raw_ing_name,
                            organization_id=org_id,
                            location_id=loc_id,
                            category=first_row.get("category") or "General",
                            base_uom=ing_row.get("unit") or "EA",
                            reporting_uom=ing_row.get("unit") or "EA",
                            current_cost=default_cost
                        )
                        db.add(matched_item)
                        db.flush()
                        items.append(matched_item)

                    existing_ing = db.query(RecipeIngredient).filter(
                        RecipeIngredient.recipe_id == recipe.id,
                        RecipeIngredient.recipe_variant_id == variant.id,
                        RecipeIngredient.inventory_item_id == matched_item.id
                    ).first()

                    if not existing_ing:
                        db.add(RecipeIngredient(
                            recipe_id=recipe.id,
                            recipe_variant_id=variant.id,
                            inventory_item_id=matched_item.id,
                            quantity=Decimal(str(ing_row.get("quantity") or 1.0)),
                            uom=ing_row.get("unit") or matched_item.base_uom or "EA",
                            notes=ing_row.get("notes"),
                            sort_order=int(ing_row.get("sort_order") or 0)
                        ))

        import_record.status = "Imported"
        db.commit()

        logger.info(f"[Recipe Import Commit Success] ID={import_id} RecipesCommitted={len(created_recipe_ids)}")
        return {
            "success": True,
            "message": f"Successfully imported {len(created_recipe_ids)} recipes into live system.",
            "recipes_committed": len(created_recipe_ids)
        }
    except Exception as exc:
        db.rollback()
        logger.error(f"[Recipe Import Commit Error] ID={import_id} Exception={exc}", exc_info=True)
        raise HTTPException(500, f"Unable to commit recipe import: {exc}")

@router.get("")
def list_recipes(db: Session = Depends(get_db)):
    recipes = db.query(Recipe).filter(Recipe.is_active == True, Recipe.location_id == current_location_id.get()).all()
    results = []
    all_items = db.query(InventoryItem).execution_options(skip_tenant_scope=True).all()
    
    for r in recipes:
        total_cost = Decimal("0.00")
        ing_details = []
        for ing in r.ingredients:
            item = db.query(InventoryItem).filter(InventoryItem.id == ing.inventory_item_id).first()
            if not item or not item.current_cost or Decimal(str(item.current_cost)) == Decimal("0.00"):
                matched = _find_item_match(item.name if item else "Ingredient", all_items)
                if matched and matched.current_cost and Decimal(str(matched.current_cost)) > 0:
                    item = matched

            if item:
                unit_cost = Decimal(str(item.current_cost or 0.0))
                if unit_cost == 0:
                    unit_cost = Decimal(str(item.last_purchase_cost or item.previous_cost or 1.50))
                factor = AvTEngine.get_unit_conversion_factor(db, item.id, ing.uom, item.base_uom)
                base_qty = Decimal(str(ing.quantity)) * factor
                cost = base_qty * unit_cost
                total_cost += cost
                ing_details.append({
                    "id": ing.id,
                    "item_id": item.id,
                    "item_name": item.name,
                    "quantity": float(ing.quantity),
                    "uom": ing.uom,
                    "unit_cost": float(unit_cost),
                    "extended_cost": float(cost)
                })

        yield_val = r.serving_yield or Decimal("1.0")
        unit_recipe_cost = total_cost / yield_val if yield_val > 0 else total_cost
        menu_price = r.menu_price or Decimal("0.00")
        food_cost_pct = float((unit_recipe_cost / menu_price * 100)) if menu_price > 0 else 0.0
        contribution_margin = menu_price - unit_recipe_cost

        results.append({
            "id": r.id,
            "name": r.name,
            "pos_identifier": r.pos_identifier,
            "category": r.category,
            "serving_yield": float(r.serving_yield),
            "menu_price": float(r.menu_price),
            "recipe_cost": float(unit_recipe_cost),
            "food_cost_pct": round(food_cost_pct, 2),
            "contribution_margin": float(contribution_margin),
            "ingredients": ing_details
        })

    return results

@router.post("")
def create_recipe(recipe_in: RecipeCreate, db: Session = Depends(get_db)):
    existing = db.query(Recipe).filter(Recipe.name == recipe_in.name, Recipe.location_id == current_location_id.get()).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"Recipe '{recipe_in.name}' already exists.")

    recipe = Recipe(
        name=recipe_in.name,
        pos_identifier=recipe_in.pos_identifier,
        category=recipe_in.category,
        serving_yield=recipe_in.serving_yield,
        menu_price=recipe_in.menu_price
        ,recipe_type=recipe_in.recipe_type, area=recipe_in.area, description=recipe_in.description
    )
    db.add(recipe)
    db.flush()

    default_variant = RecipeVariant(recipe_id=recipe.id, name="Standard", yield_quantity=recipe_in.serving_yield, yield_unit="EA")
    db.add(default_variant); db.flush()
    for ing in recipe_in.ingredients:
        if not ing.inventory_item_id and not ing.sub_recipe_id:
            raise HTTPException(400, "Each ingredient must reference an item or sub-recipe.")
        if ing.sub_recipe_id == recipe.id:
            raise HTTPException(400, "A recipe cannot reference itself.")
        r_ing = RecipeIngredient(
            recipe_id=recipe.id,
            recipe_variant_id=default_variant.id, inventory_item_id=ing.inventory_item_id, sub_recipe_id=ing.sub_recipe_id,
            quantity=ing.quantity,
            uom=ing.uom, waste_percent=ing.waste_percent, notes=ing.notes, sort_order=ing.sort_order
        )
        db.add(r_ing)

    for variant_in in recipe_in.variants:
        variant = RecipeVariant(recipe_id=recipe.id, name=variant_in.name, yield_quantity=variant_in.yield_quantity, yield_unit=variant_in.yield_unit, serving_size=variant_in.serving_size)
        db.add(variant); db.flush()
        for ing in variant_in.ingredients:
            if not ing.inventory_item_id and not ing.sub_recipe_id: raise HTTPException(400, "Each ingredient must reference an item or sub-recipe.")
            db.add(RecipeIngredient(recipe_id=recipe.id, recipe_variant_id=variant.id, inventory_item_id=ing.inventory_item_id, sub_recipe_id=ing.sub_recipe_id, quantity=ing.quantity, uom=ing.uom, waste_percent=ing.waste_percent, notes=ing.notes, sort_order=ing.sort_order))

    db.commit()
    return {"message": "Recipe created successfully", "id": recipe.id}

@router.delete("/{recipe_id}")
def delete_recipe(recipe_id: str, db: Session = Depends(get_db)):
    loc_id = current_location_id.get()
    query = db.query(Recipe).filter(Recipe.id == recipe_id)
    if loc_id:
        query = query.filter(Recipe.location_id == loc_id)
    recipe = query.first()
    if not recipe:
        raise HTTPException(404, "Recipe not found.")
    db.delete(recipe)
    db.commit()
    return {"message": "Recipe deleted successfully."}

@router.post("/bulk-delete")
def bulk_delete_recipes(payload: dict, db: Session = Depends(get_db)):
    recipe_ids = payload.get("ids", [])
    delete_all = payload.get("delete_all", False)

    loc_id = current_location_id.get()
    query = db.query(Recipe)
    if loc_id:
        query = query.filter(Recipe.location_id == loc_id)

    if not delete_all:
        if not recipe_ids:
            return {"message": "No recipes specified.", "deleted_count": 0}
        query = query.filter(Recipe.id.in_(recipe_ids))

    recipes = query.all()
    count = len(recipes)
    for r in recipes:
        db.delete(r)
    db.commit()
    return {"message": f"Successfully deleted {count} recipes.", "deleted_count": count}
